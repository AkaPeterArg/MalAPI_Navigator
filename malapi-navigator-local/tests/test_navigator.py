"""Suite funcional de MalAPI Navigator (Playwright, Chromium headless).
Cada bloque usa un contexto de navegador nuevo, así que arranca sin localStorage."""
import asyncio, json, sys, traceback
from playwright.async_api import async_playwright

from pathlib import Path
URL = (Path(__file__).resolve().parents[1] / "malapi-navigator.html").as_uri()
RESULTS = []

def check(group, name, cond, detail=""):
    RESULTS.append((group, name, bool(cond), detail))

SAMPLE_IAT = """KERNEL32.dll!OpenProcess
KERNEL32.dll!VirtualAllocEx
KERNEL32.dll!WriteProcessMemory
KERNEL32.dll!CreateRemoteThread
KERNEL32.dll!IsDebuggerPresent
KERNEL32.dll!CheckRemoteDebuggerPresent
KERNEL32.dll!GetProcAddress
KERNEL32.dll!LoadLibraryW
KERNEL32.dll!GetLastError
ADVAPI32.dll!RegOpenKeyExW
ADVAPI32.dll!RegSetValueExW
USER32.dll!GetAsyncKeyState 7"""

FAKE_CLAUDE = """
(() => {
  // La base simulada vive en el proceso de Python del test (expose_function), no en el navegador:
  // así sobrevive a recargas igual que la base real de claude.ai.
  const docs = new Map();
  const ready = window.__fakeDbLoad().then(t => { for (const [k, v] of Object.entries(JSON.parse(t || '{}'))) docs.set(k, v);
    window.__initKeys = [...docs.keys()].map(k => k.split('/').pop()); });
  const persist = () => window.__fakeDbSave(JSON.stringify(Object.fromEntries(docs)));
  window.__writes = [];
  window.__saved = [];
  const mkDoc = (path) => ({
    id: path.split('/').pop(), path,
    async get(){ await ready; const b = docs.get(path); return {id:this.id, exists:!!b, data:()=>b, metadata:{}}; },
    async set(b){ await ready; if(JSON.stringify(b).length > 256*1024) throw {code:'invalid_argument'}; docs.set(path, JSON.parse(JSON.stringify(b))); window.__writes.push(['set', path]); await persist(); },
    async update(b){ await ready; docs.set(path, Object.assign({}, docs.get(path), b)); await persist(); },
    async delete(){ await ready; docs.delete(path); window.__writes.push(['delete', path]); await persist(); }
  });
  const db = {
    doc: mkDoc,
    collection: (cp) => ({
      doc: (id) => mkDoc(cp + '/' + id),
      async get(){ await ready; const out = [];
        for (const [p, b] of docs) if (p.startsWith(cp + '/') && p.split('/').length === cp.split('/').length + 1)
          out.push({id: p.split('/').pop(), exists: true, data: () => b, metadata: {}});
        return {docs: out, size: out.length, empty: !out.length}; }
    })
  };
  const user = { id: async () => 'u123', isOwner: async () => true, canEdit: async () => true };
  const downloads = { save: async ({filename, data}) => { const text = typeof data === 'string' ? data : await data.text(); window.__saved.push({filename, text}); return {status:'saved'}; } };
  window.claude = { use: async (n) => ({db, user, downloads})[n] || null };
})();
"""

async def new_page(browser, fake=False, viewport=(1500, 950), scheme="light"):
    ctx = await browser.new_context(viewport={"width": viewport[0], "height": viewport[1]}, color_scheme=scheme)
    if fake:
        box = {"db": "{}"}
        await ctx.expose_function("__fakeDbLoad", lambda: box["db"])
        await ctx.expose_function("__fakeDbSave", lambda t: box.update(db=t))
        await ctx.add_init_script(FAKE_CLAUDE)
    pg = await ctx.new_page()
    errs = []
    pg.on("pageerror", lambda e: errs.append("pageerror: " + str(e)))
    pg.on("console", lambda m: errs.append("console: " + m.text) if m.type == "error" and "Failed to load resource" not in m.text else None)
    await pg.goto(URL)
    await pg.wait_for_timeout(400)
    return ctx, pg, errs

async def do_import(pg, text, dest="new", mode="weight", fixed=None, name=None, sample=None, sha=None):
    await pg.click(".toolbar button[data-act=importApis]")
    await pg.fill("#imText", text)
    await pg.select_option("#imDest", dest)
    await pg.select_option("#imMode", mode)
    if fixed is not None: await pg.fill("#imFixed", str(fixed))
    if name is not None and dest == "new": await pg.fill("#imName", name)
    if sample: await pg.fill("#imSample", sample)
    if sha: await pg.fill("#imHash", sha)
    await pg.click("#dlg .dlg-f button.primary")
    await pg.wait_for_timeout(150)

E = lambda pg, js: pg.evaluate(js)

async def wait_saved(pg, timeout=8000):
    await pg.wait_for_function("store.mode==='db' && store.state==='ok' && !store.busy && !store.dirty.size && !store.del.size", timeout=timeout)

# ------------------------------------------------------------------ tests
async def t_carga(b):
    g = "1. Carga y datos"
    ctx, pg, errs = await new_page(b)
    check(g, "370 APIs en el dataset", await E(pg, "API.length") == 370)
    check(g, "8 categorías en orden MalAPI", await E(pg, "CATS.join()") == "Enumeration,Injection,Evasion,Spying,Internet,Anti-Debugging,Ransomware,Helper")
    check(g, "428 asignaciones API-categoría", await E(pg, "API.reduce((s,a)=>s+a.cats.length,0)") == 428)
    check(g, "53 APIs multicategoría", await E(pg, "API.filter(a=>a.cats.length>1).length") == 53)
    check(g, "Conteos por columna = hoja Resumen", await E(pg, "CATS.map((c,i)=>API.filter(a=>a.cats.includes(i)).length).join()") == "73,91,38,29,41,24,21,111")
    check(g, "428 celdas renderizadas", await pg.locator(".cell").count() == 428)
    check(g, "Sin variantes de mayúsculas en DLL", await E(pg, "new Set(API.map(a=>a.dll.toLowerCase())).size === new Set(API.map(a=>a.dll)).size"))
    check(g, "ReadFile sin descripción ni DLL", await E(pg, "(()=>{const a=BY_NAME.get('ReadFile');return !a.desc && !a.dll})()"))
    await pg.click(".cell[data-api=ReadFile]")
    check(g, "Ficha de ReadFile avisa que falta descripción", "no incluye descripción" in await pg.inner_text("#pbody"))
    await pg.click(".cell[data-api=VirtualAllocEx]")
    check(g, "Enlace de Microsoft etiquetado como tal", await pg.locator("#pbody a", has_text="Documentación de Microsoft").count() == 1)
    await pg.click(".cell[data-api=LdrLoadDll]")
    check(g, "Enlace externo no se presenta como Microsoft", await pg.locator("#pbody a", has_text="Referencia en undocumented.ntinternals.net").count() == 1)
    check(g, "Links de documentación con URL absoluta", await E(pg, "API.filter(a=>a.doc).every(a=>/^https?:\\/\\//.test(a.doc))"))
    check(g, "Estado de guardado local", "navegador" in await pg.inner_text("#saveState"))
    check(g, "Callout de capa vacía visible", await pg.locator(".callout").count() == 1)
    check(g, "APIs de todas las combinaciones existen en el catálogo", await E(pg, "PATTERNS.flatMap(p=>p.g?p.g.flat():p.any).every(n=>BY_NAME.has(n))"))
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_matching(b):
    g = "2. Normalización de nombres"
    ctx, pg, errs = await new_page(b)
    cases = {
        "VirtualAllocExW": "VirtualAllocEx", "CreateFileW": "CreateFileA", "createfilea": "CreateFileA",
        "ZwUnmapViewOfSection": "NtUnmapViewOfSection", "_Sleep@4": "Sleep", "__imp_VirtualProtect": "VirtualProtect",
        "DnsQuery_W": "DnsQuery_A", "connect": "Connect", "LoadLibraryW": "LoadLibraryA", "LoadLibraryExW": "LoadLibraryExA",
        "RegSetValueExW": "RegSetValueExA", "SetWindowsHookExW": "SetWindowsHookExA", "CreateProcessWithTokenW": "CreateProcessWithTokenW",
        "lstrcatW": "lstrcatA", "timeGetTime": "TimeGetTime", "GetLastError": None, "ExitProcess": None, "Foo": None,
    }
    got = await E(pg, f"(()=>{{const c={json.dumps(list(cases))};return c.map(x=>{{const a=matchApi(x);return a?a.name:null}})}})()")
    for (k, v), r in zip(cases.items(), got):
        check(g, f"{k} → {v}", r == v, f"obtuvo {r}")
    await ctx.close()

async def t_parser(b):
    g = "3. Parser de importaciones"
    ctx, pg, errs = await new_page(b)
    formats = {
        "una por línea": ("VirtualAllocEx\nWriteProcessMemory\nCreateRemoteThread", 3, 0),
        "dll!func": ("KERNEL32.dll!VirtualAllocEx\nkernel32.dll!WriteProcessMemory", 2, 0),
        "dll.func": ("KERNEL32.VirtualAllocEx\nADVAPI32.RegSetValueExW", 2, 0),
        "pefile": ("  0x40a000 VirtualAllocEx\n  0x40a004 GetLastError\n  0x40a008 ExitProcess", 1, 2),
        "dumpbin": ("    KERNEL32.dll\n                 1A0 VirtualAllocEx\n                 2B1 GetModuleHandleW", 2, 0),
        "rabin2 -i": ("1   0x00402000 NONE FUNC KERNEL32.dll OpenProcess\n2   0x00402004 NONE FUNC KERNEL32.dll CloseHandle", 1, 1),
        "CSV PEStudio": ('"VirtualAllocEx","kernel32.dll","x"\n"HeapFree","kernel32.dll",""', 1, 1),
        "comentarios #": ("# encabezado VirtualAlloc\nSleep", 1, 0),
        "vacío": ("", 0, 0),
    }
    for name, (txt, nf, nu) in formats.items():
        r = await E(pg, f"(()=>{{const r=parseImports({json.dumps(txt)});return [r.found.size,r.unmatched.length,r.unmatched]}})()")
        check(g, f"Formato {name}: {nf} coincidencias, {nu} fuera de catálogo", r[0] == nf and r[1] == nu, f"obtuvo {r}")
    r = await E(pg, "(()=>{const r=parseImports('VirtualAllocEx 7\\nSleep,2.5\\nCreateRemoteThread');return [...r.found.values()].map(x=>[x.a.name,x.score])})()")
    check(g, "Score explícito al final de línea (incluye decimales)", r == [["VirtualAllocEx", 7], ["Sleep", 2.5], ["CreateRemoteThread", None]], r)
    r = await E(pg, "(()=>{const r=parseImports('VirtualAllocEx\\nVirtualAllocExW\\nvirtualallocex');return [r.found.size,[...r.found.get('VirtualAllocEx').as]]})()")
    check(g, "Duplicados y variantes se consolidan", r[0] == 1 and "VirtualAllocExW" in r[1], r)
    await ctx.close()

async def t_import_ui(b):
    g = "4. Importación desde la interfaz"
    ctx, pg, errs = await new_page(b)
    await pg.click(".toolbar button[data-act=importApis]")
    await pg.click("#dlg .dlg-f button.primary")
    check(g, "Texto sin APIs: muestra error y no cierra", await E(pg, "document.getElementById('dlg').open") and "No se encontró" in await pg.inner_text("#imPrev"))
    await pg.fill("#imText", SAMPLE_IAT)
    check(g, "Vista previa en vivo", "Coinciden 11" in await pg.inner_text("#imPrev"), await pg.inner_text("#imPrev"))
    await pg.click("#dlg [data-close]")
    await do_import(pg, SAMPLE_IAT, name="Dropper X", sample="invoice.exe", sha="ab" * 32)
    check(g, "Crea capa nueva y la activa", await E(pg, "L().name") == "Dropper X" and await E(pg, "S.order.length") == 2)
    check(g, "11 APIs anotadas", await E(pg, "Object.keys(L().entries).length") == 11)
    check(g, "Score por peso de categoría (VirtualAllocEx=5)", await E(pg, "L().entries.VirtualAllocEx.s") == 5)
    check(g, "Score por peso Helper (RegOpenKeyExA=1)", await E(pg, "L().entries.RegOpenKeyExA.s") == 1)
    check(g, "Score explícito respetado (GetAsyncKeyState=7)", await E(pg, "L().entries.GetAsyncKeyState.s") == 7)
    check(g, "Comentario 'Importada como' para variantes", "LoadLibraryW" in (await E(pg, "L().entries.LoadLibraryA.m") or ""))
    check(g, "Fuera de catálogo guardadas", await E(pg, "L().unmatched") == ["GetLastError"])
    check(g, "Muestra y SHA-256 guardados", await E(pg, "L().sample.name+'|'+L().sample.hash.length") == "invoice.exe|64")
    check(g, "Gradiente ajustado al rango (0–7)", await E(pg, "[L().gradient.min,L().gradient.max]") == [0, 7])
    check(g, "Abre el Perfil al terminar", await E(pg, "UI.ptab") == "profile")
    check(g, "Contador de la pestaña = 11", (await pg.inner_text(".tab[aria-selected=true] .tcount")) == "11")
    await do_import(pg, "Sleep\nVirtualAllocEx", dest="active", mode="fixed", fixed=2)
    check(g, "Fusión en capa activa sin crear capa", await E(pg, "S.order.length") == 2 and await E(pg, "Object.keys(L().entries).length") == 12)
    check(g, "Valor fijo aplicado en fusión", await E(pg, "L().entries.Sleep.s") == 2)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_seleccion_edicion(b):
    g = "5. Selección y edición"
    ctx, pg, errs = await new_page(b)
    await pg.click(".cell[data-api=VirtualAllocEx]")
    check(g, "Clic abre la ficha", (await pg.inner_text("#pbody h2")) == "VirtualAllocEx")
    check(g, "Ficha con DLL, categorías y enlaces", await pg.locator("#pbody a[href*='malapi.io/winapi/VirtualAllocEx']").count() == 1 and "Kernel32.dll" in await pg.inner_text("#pbody"))
    await pg.fill("#edScore", "3.5"); await pg.press("#edScore", "Tab")
    check(g, "Score decimal desde input", await E(pg, "L().entries.VirtualAllocEx.s") == 3.5)
    check(g, "Badge de score en la celda", (await pg.inner_text(".cell[data-api=VirtualAllocEx] .sc")) == "3,5")
    await pg.click("#pbody .quick button[data-v='4']")
    check(g, "Score rápido", await E(pg, "L().entries.VirtualAllocEx.s") == 4)
    await pg.click("#pbody .sw[data-v='#5aa469']")
    check(g, "Color manual", await E(pg, "L().entries.VirtualAllocEx.c") == "#5aa469")
    bg = await E(pg, "getComputedStyle(document.querySelector('.cell[data-api=VirtualAllocEx]')).backgroundColor")
    check(g, "Color manual tiene prioridad sobre gradiente", bg == "rgb(90, 164, 105)", bg)
    await pg.fill("#edComment", "RVA 0x1234, regla CAPA"); await pg.press("#edComment", "Tab")
    check(g, "Comentario y marca visual", await E(pg, "L().entries.VirtualAllocEx.m") == "RVA 0x1234, regla CAPA" and await pg.locator(".cell.hascomment[data-api=VirtualAllocEx]").count() == 1)
    await pg.click("#pbody .sw.none")
    check(g, "Quitar color vuelve al gradiente", await E(pg, "L().entries.VirtualAllocEx.c") is None)
    # multi-api: aparece en todas sus columnas
    await pg.click(".cell[data-api=GetProcAddress] >> nth=0")
    check(g, "API multicategoría seleccionada en todas sus columnas", await pg.locator(".cell.sel[data-api=GetProcAddress]").count() == 2)
    await pg.click(".cell[data-api=OpenProcess]", modifiers=["Control"])
    check(g, "Ctrl+clic suma a la selección", await E(pg, "UI.sel.size") == 2)
    await pg.click(".cell[data-api=OpenProcess]", modifiers=["Control"])
    check(g, "Ctrl+clic de nuevo la quita", await E(pg, "UI.sel.size") == 1)
    first = await E(pg, "UI.colLists[1][0]"); fifth = await E(pg, "UI.colLists[1][4]")
    await pg.click(f".cell[data-col='1'][data-api='{first}']")
    await pg.click(f".cell[data-col='1'][data-api='{fifth}']", modifiers=["Shift"])
    check(g, "Mayús+clic selecciona rango de 5", await E(pg, "UI.sel.size") == 5)
    await pg.click("#pbody .quick button[data-v='2']")
    check(g, "Edición en lote aplica a las 5", await E(pg, "[...UI.sel].every(n=>L().entries[n]&&L().entries[n].s===2)"))
    await pg.keyboard.press("Escape")
    check(g, "Esc deselecciona", await E(pg, "UI.sel.size") == 0)
    await pg.click(".cell[data-api=Sleep] >> nth=0")
    await pg.keyboard.press("7")
    check(g, "Atajo numérico asigna score", await E(pg, "L().entries.Sleep.s") == 7)
    await pg.keyboard.press("Delete")
    check(g, "Supr borra la anotación", await E(pg, "L().entries.Sleep") is None)
    await pg.fill("#edComment", "tecla 5 no debe cambiar score")
    await pg.keyboard.press("5")
    check(g, "Atajos no interfieren al escribir", await E(pg, "L().entries.Sleep") is None)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_filtros(b):
    g = "6. Búsqueda, filtros y orden"
    ctx, pg, errs = await new_page(b)
    await do_import(pg, SAMPLE_IAT)
    await pg.fill("#q", "keylogger"); await pg.wait_for_timeout(250)
    hits = await pg.locator(".cell.hit").count()
    check(g, "Búsqueda por descripción resalta coincidencias", hits > 0 and await pg.locator(".cell.dim").count() > 0, hits)
    check(g, "Botón muestra conteo", f"({hits})" in await pg.inner_text("#btnSelHits"))
    await pg.click("#btnSelHits")
    check(g, "Seleccionar coincidencias", await E(pg, "UI.sel.size") > 0 and await E(pg, "[...UI.sel].every(n=>matches(BY_NAME.get(n),'keylogger'))"))
    await pg.fill("#q", ""); await pg.wait_for_timeout(250)
    check(g, "Limpiar búsqueda quita atenuado", await pg.locator(".cell.dim").count() == 0)
    await pg.select_option("#dll", "wininet")
    check(g, "Filtro por DLL", await E(pg, "[...document.querySelectorAll('.cell')].every(c=>BY_NAME.get(c.dataset.api).dll==='Wininet.dll')") and await pg.locator(".cell").count() > 0)
    await pg.select_option("#dll", "ws2_32")
    check(g, "Filtro ws2_32 agrupa .dll y .lib", await E(pg, "new Set([...document.querySelectorAll('.cell')].map(c=>BY_NAME.get(c.dataset.api).dll)).size") == 2)
    await pg.select_option("#dll", "")
    await pg.select_option("#show", "marked")
    check(g, "Mostrar solo anotadas", await pg.locator(".cell").count() == await E(pg, "API.filter(a=>L().entries[a.name]).reduce((s,a)=>s+a.cats.length,0)"))
    await pg.select_option("#show", "unmarked")
    check(g, "Mostrar sin anotar", await E(pg, "[...document.querySelectorAll('.cell')].every(c=>!L().entries[c.dataset.api])"))
    await pg.select_option("#show", "all")
    await pg.select_option("#sort", "score")
    col = await E(pg, "UI.colLists[1].slice(0,7).map(n=>L().entries[n]?L().entries[n].s:null)")
    check(g, "Orden por score descendente", col[:6] == sorted(col[:6], reverse=True) and col[0] == 5, col)
    await pg.select_option("#sort", "name")
    names = await E(pg, "UI.colLists[0]")
    check(g, "Orden alfabético", names == sorted(names, key=lambda s: s.lower()) or names == sorted(names), names[:5])
    await pg.select_option("#sort", "weight")
    w = await E(pg, "UI.colLists[0].map(n=>apiWeight(BY_NAME.get(n)))")
    check(g, "Orden por peso intrínseco", w == sorted(w, reverse=True))
    await pg.select_option("#sort", "dll")
    d = await E(pg, "UI.colLists[1].map(n=>BY_NAME.get(n).dll)")
    check(g, "Orden por DLL", d == sorted(d))
    await pg.check("#colsByWeight")
    heads = await pg.locator(".ch-name").all_inner_texts()
    sums = await E(pg, "columns().map(c=>c.sum)")
    check(g, "Columnas por peso (Σ descendente)", sums == sorted(sums, reverse=True) and heads[0] == "Injection", heads)
    await pg.check("#dense")
    check(g, "Vista compacta", await pg.locator(".matrix.dense").count() == 1)
    await pg.click("button[data-act=columns]")
    await pg.uncheck("#dlg input[value='7']"); await pg.click("#dlg .dlg-f button.primary")
    check(g, "Ocultar columna Helper", "Helper" not in await pg.locator(".ch-name").all_inner_texts() and await pg.locator(".col").count() == 7)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_capa(b):
    g = "7. Gestión de capas y gradiente"
    ctx, pg, errs = await new_page(b)
    await do_import(pg, SAMPLE_IAT, name="Original")
    await pg.click("[data-ptab=layer]")
    await pg.fill("#lyName", "Renombrada"); await pg.press("#lyName", "Tab")
    check(g, "Renombrar actualiza pestaña", "Renombrada" in await pg.inner_text(".tab[aria-selected=true]"))
    await pg.fill("#gMax", "10"); await pg.press("#gMax", "Tab")
    check(g, "Cambiar máximo del gradiente", await E(pg, "L().gradient.max") == 10)
    await pg.click("button[data-act=fitGradient]")
    check(g, "Ajustar al rango", await E(pg, "[L().gradient.min,L().gradient.max]") == [0, 7])
    await pg.click("button[data-act=gradPreset][data-v='3']")
    check(g, "Preset de gradiente", await E(pg, "L().gradient.colors[2]") == "#08519c")
    c = await E(pg, "getComputedStyle(document.querySelector('.cell[data-api=GetAsyncKeyState]')).backgroundColor")
    check(g, "Score máximo pinta color alto", c == "rgb(8, 81, 156)", c)
    await pg.click("button[data-act=duplicate]")
    check(g, "Duplicar", await E(pg, "S.order.length") == 3 and "(copia)" in await E(pg, "L().name") and await E(pg, "Object.keys(L().entries).length") == 11)
    await E(pg, "L().entries.Sleep={s:9}")
    check(g, "Duplicado es independiente", await E(pg, "S.layers[S.order[1]].entries.Sleep") is None)
    await pg.click("[data-ptab=layer]")
    await pg.click("button[data-act=moveLeft]")
    check(g, "Mover capa", await E(pg, "S.order.indexOf(S.activeId)") == 1)
    await pg.click("button[data-act=clearLayer]")
    check(g, "Vaciar anotaciones", await E(pg, "Object.keys(L().entries).length") == 0)
    await pg.click("[data-ptab=layer]")
    await pg.click("button[data-act=deleteLayer]"); await pg.click("#dlg .dlg-f button.danger")
    check(g, "Eliminar capa", await E(pg, "S.order.length") == 2)
    await pg.click(".tab[data-layer] >> nth=0")
    check(g, "Cambiar de capa con pestañas", await E(pg, "S.activeId===S.order[0]"))
    await E(pg, "S.order=[S.order[0]];"); await E(pg, "renderPanel()")
    await pg.click("[data-ptab=layer]")
    check(g, "No permite eliminar la única capa", await pg.locator("button[data-act=deleteLayer]").count() == 0)
    # XSS
    await pg.fill("#lyName", "<img src=x onerror=window.__xss=1>"); await pg.press("#lyName", "Tab")
    await pg.click("[data-ptab=profile]"); await pg.click(".tab[aria-selected=true]")
    await pg.wait_for_timeout(100)
    check(g, "Nombres con HTML se escapan (sin XSS)", await E(pg, "window.__xss") is None and await pg.locator(".tab img").count() == 0)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_expr(b):
    g = "8. Motor de expresiones"
    ctx, pg, errs = await new_page(b)
    cases = [
        ("a + b", {"a": 2, "b": 3}, 5), ("a - b * 2", {"a": 10, "b": 3}, 4), ("(a - b) * 2", {"a": 10, "b": 3}, 14),
        ("max(a, b, c)", {"a": 1, "b": 7, "c": 3}, 7), ("min(a, b)", {"a": 1, "b": 0}, 0), ("avg(a, b)", {"a": 2, "b": 5}, 3.5),
        ("abs(a - b)", {"a": 1, "b": 4}, 3), ("round(a / b)", {"a": 7, "b": 2}, 4), ("a / b", {"a": 5, "b": 0}, 0),
        ("a > 0 && b > 0 ? a + b : 0", {"a": 2, "b": 0}, 0), ("a > 0 && b > 0 ? a + b : 0", {"a": 2, "b": 1}, 3),
        ("a > 0 || b > 0", {"a": 0, "b": 1}, 1), ("!a", {"a": 0}, 1), ("-a + 1", {"a": 3}, -2), ("a % 3", {"a": 7}, 1),
        ("a == 2 ? 10 : 20", {"a": 2}, 10), ("a != 2", {"a": 2}, 0), ("a >= 2", {"a": 2}, 1), ("a <= 1", {"a": 2}, 0),
        ("a", {}, 0), ("1.5 * a", {"a": 2}, 3), ("A + B", {"a": 1, "b": 2}, 3),
    ]
    for expr, env, want in cases:
        got = await E(pg, f"compileExpr({json.dumps(expr)}).fn({json.dumps(env)})")
        check(g, f"{expr} con {env} = {want}", abs(got - want) < 1e-9, got)
    for bad, frag in [("a +", "termina"), ("a + $", "no válido"), ("foo(a)", "Función desconocida"), ("ab + 1", "Variable desconocida"), ("(a + b", "Se esperaba"), ("a b", "Sobra"), ("", "Escribí"), ("1..2", "Número inválido")]:
        msg = await E(pg, f"(()=>{{try{{compileExpr({json.dumps(bad)});return 'OK'}}catch(e){{return e.message}}}})()")
        check(g, f"Rechaza «{bad}»", frag in msg, msg)
    check(g, "No usa eval ni Function (compatible con CSP)", await E(pg, "!/\\beval\\(|new Function/.test(document.scripts[0].textContent)"))
    await ctx.close()

async def t_overlay(b):
    g = "9. Superposición de capas"
    ctx, pg, errs = await new_page(b)
    await pg.click(".toolbar button[data-act=overlay]")
    check(g, "Con una sola capa avisa", "al menos dos capas" in await pg.inner_text("#dlg"))
    await pg.click("#dlg .dlg-f button.primary")
    await do_import(pg, "VirtualAllocEx 4\nWriteProcessMemory 2\nSleep 3", name="A")
    await pg.click(".cell[data-api=VirtualAllocEx]"); await pg.click("#pbody .sw[data-v='#7a5cc0']")
    await pg.fill("#edComment", "nota A"); await pg.press("#edComment", "Tab")
    await do_import(pg, "VirtualAllocEx 1\nCreateRemoteThread 5", name="B")
    await pg.click(".cell[data-api=VirtualAllocEx]"); await pg.fill("#edComment", "nota B"); await pg.press("#edComment", "Tab")
    await pg.click(".toolbar button[data-act=overlay]")
    boxes = pg.locator("#ovList input")
    await boxes.nth(0).uncheck(); await boxes.nth(1).check(); await boxes.nth(2).check()
    check(g, "Letras asignadas en orden (a=A, b=B)", await pg.locator("#ovList .letter.on").all_inner_texts() == ["a", "b"])
    preview = await pg.inner_text("#ovPrev")
    check(g, "Vista previa suma: 4 APIs, rango 2–5", "4 APIs" in preview and "2 a 5" in preview, preview)
    expected = {
        0: {"VirtualAllocEx": 5, "WriteProcessMemory": 2, "Sleep": 3, "CreateRemoteThread": 5},
        1: {"VirtualAllocEx": 4, "WriteProcessMemory": 2, "Sleep": 3, "CreateRemoteThread": 5},
        2: {"VirtualAllocEx": 2.5, "WriteProcessMemory": 1, "Sleep": 1.5, "CreateRemoteThread": 2.5},
        3: {"VirtualAllocEx": 1},
        4: {"VirtualAllocEx": 5},
        5: {"VirtualAllocEx": 1, "WriteProcessMemory": 1, "Sleep": 1, "CreateRemoteThread": 1},
        6: {"WriteProcessMemory": 2, "Sleep": 3},
        7: {"VirtualAllocEx": 3, "WriteProcessMemory": 2, "Sleep": 3, "CreateRemoteThread": -5},
        8: {"VirtualAllocEx": 4},
    }
    names = ["Suma", "Máximo", "Promedio", "Mínimo", "Intersección", "Unión", "Solo en a", "Diferencia", "Ponderar", "Ausentes"]
    for i, want in expected.items():
        await pg.click(f"#ovPresets button[data-i='{i}']")
        got = await E(pg, "Object.fromEntries(Object.entries(computeOverlay().entries).map(([k,v])=>[k,v.s]))")
        check(g, f"Preset {names[i]}", got == want, f"{await pg.input_value('#ovExpr')} → {got}")
    await pg.click("#ovPresets button[data-i='9']")
    n_abs = await E(pg, "computeOverlay().count")
    check(g, "Preset Ausentes en todas (370 − 4 = 366)", n_abs == 366, n_abs)
    await pg.fill("#ovExpr", "a + c")
    check(g, "Variable sin capa asignada da error", "solo hay capas asignadas" in await pg.inner_text("#ovPrev"))
    await pg.click("#dlg .dlg-f button.primary")
    check(g, "No crea capa con expresión inválida", await E(pg, "document.getElementById('dlg').open") and await E(pg, "S.order.length") == 3)
    await pg.fill("#ovExpr", "a + b")
    await pg.select_option("#ovColor", index=1)
    await pg.fill("#ovName", "Combinada")
    await pg.click("#dlg .dlg-f button.primary"); await pg.wait_for_timeout(150)
    check(g, "Crea capa resultado y la activa", await E(pg, "L().name") == "Combinada" and await E(pg, "L().source") == "overlay")
    check(g, "Hereda color manual de la capa elegida", await E(pg, "L().entries.VirtualAllocEx.c") == "#7a5cc0")
    m = await E(pg, "L().entries.VirtualAllocEx.m") or ""
    check(g, "Combina comentarios con origen", "[A] nota A" in m and "[B] nota B" in m, m)
    check(g, "Descripción documenta la expresión", "a + b" in await E(pg, "L().desc"))
    check(g, "Gradiente ajustado al resultado", await E(pg, "L().gradient.max") == 5)
    check(g, "Capas de origen intactas", await E(pg, "S.layers[S.order[1]].entries.VirtualAllocEx.s") == 4)
    # comparación en ficha
    await pg.click(".cell[data-api=VirtualAllocEx]")
    rows = await pg.locator("#pbody table.t tr").count()
    check(g, "Ficha muestra score en las otras capas", rows == 4, rows)
    # referencia x muestra
    await pg.click(".tab.add"); await pg.click(".choice button[data-c=reference]")
    check(g, "Capa de referencia: 370 APIs con su peso", await E(pg, "Object.keys(L().entries).length") == 370 and await E(pg, "L().entries.Accept.s") == 3)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_overlay_bordes(b):
    g = "9b. Superposición: casos borde"
    ctx, pg, errs = await new_page(b)
    await E(pg, """(()=>{const A=newLayer('A'),B=newLayer('B'),C=newLayer('C');
      A.entries={VirtualAllocEx:{s:4},Sleep:{c:'#d6453d'},OpenProcess:{c:'#d6453d',m:'mapeo A'}};
      B.entries={VirtualAllocEx:{s:1},OpenProcess:{c:'#3f8fbf'}};
      C.entries={Sleep:{s:2}};
      S.layers={};S.order=[];addLayer(A);addLayer(B);addLayer(C,false);S.activeId=A.id;renderAll();})()""")
    R = lambda recipe: E(pg, f"(()=>{{const r=overlayCompute({json.dumps(recipe)});return r.error?{{error:r.error}}:{{count:r.count,e:r.entries}}}})()")
    ids = await E(pg, "S.order.slice()")
    A, B, C = ids
    base = {"sources": [["a", A], ["b", B]], "colorFrom": "", "comments": False, "colorOnly": 1}
    r = await R({**base, "expr": "!a"})
    check(g, "Se evalúa sobre las 370 APIs: !a devuelve 367", r["count"] == 367, r.get("count"))
    r = await R({**base, "expr": "5 - a"})
    check(g, "Constante menos capa cubre todo el catálogo (370)", r["count"] == 370 and r["e"]["VirtualAllocEx"]["s"] == 1 and r["e"]["Accept"]["s"] == 5)
    r = await R({**base, "expr": "(a != 0 && b != 0) ? a + b : 0"})
    check(g, "Intersección incluye APIs mapeadas solo por color", set(r["e"]) == {"VirtualAllocEx", "OpenProcess"} and r["e"]["OpenProcess"]["s"] == 2, r["e"])
    r = await R({**base, "expr": "(a != 0 && b == 0) ? a : 0"})
    check(g, "Solo en a incluye mapeo por color", set(r["e"]) == {"Sleep"}, r["e"])
    r = await R({**base, "colorOnly": 0, "expr": "(a != 0 && b != 0) ? a + b : 0"})
    check(g, "Con valor 0, el color solo no cuenta", set(r["e"]) == {"VirtualAllocEx"}, r["e"])
    r = await R({**base, "colorOnly": 3, "expr": "a"})
    check(g, "Valor configurable para color solo (3)", r["e"]["Sleep"]["s"] == 3 and r["e"]["VirtualAllocEx"]["s"] == 4)
    r = await R({**base, "expr": "a + b", "colorFrom": B})
    check(g, "Color heredado de la capa elegida (b)", r["e"]["OpenProcess"].get("c") == "#3f8fbf" and "c" not in r["e"]["Sleep"])
    r = await R({**base, "expr": "a + b", "comments": True})
    check(g, "Comentarios con prefijo de capa", r["e"]["OpenProcess"].get("m") == "[A] mapeo A")
    r = await R({**base, "expr": "a / 0"})
    check(g, "División por cero da 0 (sin NaN/Infinity)", r["count"] == 0)
    r = await R({**base, "expr": "max()"})
    check(g, "max() sin argumentos no rompe", r["count"] == 0)
    r = await R({"sources": [], "expr": "1", "colorFrom": "", "comments": False, "colorOnly": 1})
    check(g, "Sin capas: error claro", "al menos una capa" in r.get("error", ""))
    # letras según el orden de la lista
    await pg.click(".toolbar button[data-act=overlay]")
    boxes = pg.locator("#ovList input")
    await boxes.nth(0).uncheck(); await boxes.nth(1).check(); await boxes.nth(2).check()
    check(g, "Letras siguen el orden de la lista (B=a, C=b)", await E(pg, "overlayRecipe().sources.map(x=>S.layers[x[1]].name).join()") == "B,C")
    await pg.fill("#ovExpr", "a + b")
    await pg.fill("#ovColorOnly", "0")
    p0 = await pg.inner_text("#ovPrev")
    await pg.fill("#ovColorOnly", "1")
    p1 = await pg.inner_text("#ovPrev")
    check(g, "Vista previa reacciona al valor de color solo", p0 != p1, [p0, p1])
    await pg.check("#ovComments"); await pg.uncheck("#ovComments")
    check(g, "Vista previa sigue válida al tocar comentarios", "Resultado" in await pg.inner_text("#ovPrev"))
    await boxes.nth(0).check(); await boxes.nth(1).uncheck(); await boxes.nth(2).uncheck()
    await pg.click("#ovPresets button[data-i='0']")
    await pg.fill("#ovName", "Sobre A+B")
    await boxes.nth(1).check()
    await pg.click("#ovPresets button[data-i='0']")
    await pg.click("#dlg .dlg-f button.primary"); await pg.wait_for_timeout(150)
    oid = await E(pg, "S.activeId")
    check(g, "Guarda la receta en la capa", await E(pg, "L().recipe && L().recipe.expr") == "a + b")
    check(g, "Resultado inicial: OpenProcess=2", await E(pg, "L().entries.OpenProcess.s") == 2)
    # editar origen y recalcular
    await E(pg, f"S.layers['{B}'].entries.Sleep={{s:10}}; touch(S.layers['{B}'])")
    check(g, "La capa no cambia sola al editar el origen", await E(pg, "L().entries.Sleep.s") == 1)
    await pg.click("[data-ptab=layer]")
    check(g, "Panel Capa muestra la receta", "a + b" in await pg.inner_text("#pbody") and await pg.locator("button[data-act=recalc]").is_enabled())
    await pg.click("button[data-act=recalc]"); await pg.wait_for_timeout(100)
    check(g, "Recalcular toma los cambios del origen (Sleep=11)", await E(pg, "L().entries.Sleep.s") == 11)
    check(g, "Recalcular reajusta el gradiente", await E(pg, "L().gradient.max") == 11)
    # persistencia de receta
    await pg.reload(); await pg.wait_for_timeout(300)
    check(g, "Receta sobrevive a recargar", await E(pg, f"S.layers['{oid}'].recipe.sources.length") == 2)
    await E(pg, f"S.activeId='{oid}'; renderAll()")
    await pg.click("[data-ptab=layer]"); await pg.click("button[data-act=duplicate]")
    check(g, "Duplicar conserva la receta", await E(pg, "!!L().recipe"))
    # origen eliminado
    await E(pg, f"deleteLayer('{B}'); S.activeId='{oid}'; renderAll()")
    await pg.click("[data-ptab=layer]")
    check(g, "Origen eliminado: botón deshabilitado y aviso", not await pg.locator("button[data-act=recalc]").is_enabled() and "capa eliminada" in await pg.inner_text("#pbody"))
    r = await E(pg, "overlayCompute(L().recipe).error || ''")
    check(g, "Origen eliminado: error explícito, sin romper", "se eliminó" in r, r)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_perfil(b):
    g = "10. Perfil y combinaciones"
    ctx, pg, errs = await new_page(b)
    await pg.click("[data-ptab=profile]")
    check(g, "Perfil vacío con llamada a importar", await pg.locator("#pbody button[data-act=importApis]").count() == 1)
    tests = {
        "Inyección remota por hilo": "OpenProcess\nVirtualAllocEx\nWriteProcessMemory\nCreateRemoteThread",
        "Process hollowing": "CreateProcessW\nZwUnmapViewOfSection\nWriteProcessMemory\nSetThreadContext\nResumeThread",
        "Keylogger por hook": "SetWindowsHookExW\nCallNextHookEx\nGetMessageW",
        "Captura de pantalla": "GetDC\nBitBlt",
        "Cifrado masivo de archivos": "FindFirstFileW\nFindNextFileW\nCryptEncrypt\nWriteFile",
        "Chequeos anti-depuración": "IsDebuggerPresent\nCheckRemoteDebuggerPresent",
        "Evasión por tiempo o sandbox": "GetTickCount\nSleep\nQueryPerformanceCounter",
        "Persistencia como servicio": "OpenSCManagerW\nCreateServiceW\nStartServiceW",
    }
    for pat, txt in tests.items():
        await do_import(pg, txt, name=pat)
        full = await E(pg, "evalPatterns(new Set(API.filter(a=>isMarked(L().entries[a.name])).map(a=>a.name))).filter(x=>x.full).map(x=>x.p.n)")
        check(g, f"Detecta: {pat}", pat in full, full)
    await do_import(pg, "OpenProcess\nVirtualAllocEx\nWriteProcessMemory", name="parcial")
    part = await E(pg, "evalPatterns(new Set(Object.keys(L().entries))).filter(x=>!x.full).map(x=>[x.p.n,x.missing])")
    check(g, "Parcial indica lo que falta", any(p[0] == "Inyección remota por hilo" and "CreateRemoteThread" in p[1][0] for p in part), part)
    await do_import(pg, "IsDebuggerPresent", name="uno")
    full = await E(pg, "evalPatterns(new Set(Object.keys(L().entries))).filter(x=>x.full).length")
    check(g, "Una sola API anti-debug no dispara el patrón k=2", full == 0)
    await do_import(pg, SAMPLE_IAT, name="perfil")
    await pg.click("[data-ptab=profile]")
    txt = await pg.inner_text("#pbody")
    check(g, "Resumen con APIs marcadas y score", "APIs marcadas" in txt and "score total" in txt)
    check(g, "Categorías ordenadas por peso", (await pg.locator("#pbody .prow-top b").all_inner_texts())[0] == "Injection")
    check(g, "Aviso de resolución dinámica", "resolución dinámica" in txt)
    check(g, "Lista fuera de catálogo", "GetLastError" in txt)
    await pg.click("#pbody button.chip.api[data-v=OpenProcess] >> nth=0")
    check(g, "Chip del perfil lleva a la ficha", await E(pg, "UI.ptab") == "sel" and (await pg.inner_text("#pbody h2")) == "OpenProcess")
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_pesos(b):
    g = "11. Pesos por categoría"
    ctx, pg, errs = await new_page(b)
    await pg.click("button[data-act=weights]")
    await pg.fill("#w-Helper", "3"); await pg.click("#dlg .dlg-f button.primary")
    check(g, "Guardar pesos", await E(pg, "S.weights.Helper") == 3)
    check(g, "Peso intrínseco usa el máximo (GetDriveTypeA)", await E(pg, "apiWeight(BY_NAME.get('GetDriveTypeA'))") == 5)
    await do_import(pg, "RegOpenKeyExW")
    check(g, "Importación usa los pesos nuevos", await E(pg, "L().entries.RegOpenKeyExA.s") == 3)
    await pg.click("button[data-act=weights]")
    await pg.click("#dlg .dlg-f button >> text=Restablecer")
    check(g, "Restablecer rellena valores por defecto", await pg.input_value("#w-Helper") == "1")
    await pg.click("#dlg .dlg-f button.primary")
    check(g, "Capas existentes no cambian", await E(pg, "L().entries.RegOpenKeyExA.s") == 3)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_export_import(b):
    g = "12. Exportar e importar"
    ctx, pg, errs = await new_page(b)
    await do_import(pg, SAMPLE_IAT, name="Exp", sample="s.exe", sha="cd" * 32)
    await pg.click(".cell[data-api=VirtualAllocEx]"); await pg.click("#pbody .sw[data-v='#d6453d']")
    await pg.fill("#edComment", 'con "comillas", y coma'); await pg.press("#edComment", "Tab")
    async def dl(act):
        await pg.click("[data-ptab=layer]")
        async with pg.expect_download() as info:
            await pg.click(f"button[data-act={act}]")
        d = await info.value
        return d.suggested_filename, open(await d.path(), encoding="utf-8", newline="").read()
    fn, txt = await dl("exportLayer")
    check(g, "Local: la exportación descarga un archivo real", fn == "malapi-capa-Exp.json", fn)
    exported = json.loads(txt)
    check(g, "JSON con formato y metadatos", exported["format"] == "malapi-layer" and exported["sample"]["name"] == "s.exe" and len(exported["entries"]) == 11)
    e = next(x for x in exported["entries"] if x["api"] == "VirtualAllocEx")
    check(g, "Entrada completa (score, color, comentario, categorías)", e["score"] == 5 and e["color"] == "#d6453d" and e["categories"] == ["Injection"] and "comillas" in e["comment"])
    fn, csv = await dl("exportCsv")
    lines = csv.lstrip("\ufeff").split("\r\n")
    check(g, "CSV descargado con BOM, CRLF, encabezado y 11 filas", csv.startswith("\ufeff") and lines[0].startswith("API,DLL") and len(lines) == 12 and fn.endswith(".csv"))
    check(g, "CSV escapa comillas", '"con ""comillas"", y coma"' in csv)
    check(g, "CSV ordenado por score", lines[1].split(",")[0].strip('"') == "GetAsyncKeyState")
    fn, wtxt = await dl("exportAll")
    ws = json.loads(wtxt)
    check(g, "Exportar todas: workspace con 2 capas", ws["format"] == "malapi-workspace" and len(ws["layers"]) == 2)
    # reimportar (roundtrip)
    await E(pg, f"importJsonText({json.dumps(json.dumps(exported))})")
    check(g, "Reimportar JSON propio", await E(pg, "L().name") == "Exp" and await E(pg, "L().source") == "json" and await E(pg, "S.order.length") == 3)
    check(g, "Roundtrip conserva score, color y comentario", await E(pg, "JSON.stringify(L().entries.VirtualAllocEx)") == json.dumps({"s": 5, "c": "#d6453d", "m": 'con "comillas", y coma'}, separators=(",", ":"), ensure_ascii=False))
    check(g, "Roundtrip conserva muestra y fuera de catálogo", await E(pg, "L().sample.hash.length") == 64 and await E(pg, "L().unmatched") == ["GetLastError"])
    await E(pg, f"importJsonText({json.dumps(json.dumps(ws))})")
    check(g, "Importar workspace completo", await E(pg, "S.order.length") == 5)
    nav = {"name": "Navigator", "gradient": {"colors": ["#ff6666ff", "#ffe766ff", "#8ec843ff"], "minValue": 0, "maxValue": 100},
           "techniques": [{"techniqueID": "VirtualAllocExW", "score": 80, "color": "#e60d0d", "comment": "x"}, {"techniqueID": "T1055", "score": 50}]}
    await E(pg, f"importJsonText({json.dumps(json.dumps(nav))})")
    check(g, "Formato ATT&CK Navigator (techniques)", await E(pg, "L().entries.VirtualAllocEx.s") == 80 and await E(pg, "L().gradient.max") == 100)
    check(g, "Colores RGBA de Navigator recortados", await E(pg, "L().gradient.colors[0]") == "#ff6666")
    check(g, "IDs no reconocidos van a fuera de catálogo", await E(pg, "L().unmatched") == ["T1055"])
    n = await E(pg, "S.order.length")
    await E(pg, "importJsonText('no es json')")
    check(g, "JSON inválido no rompe ni crea capas", await E(pg, "S.order.length") == n and "no es un JSON" in await pg.inner_text("#toast"))
    await E(pg, "importJsonText(JSON.stringify({entries:{NoExiste:{s:1}, Sleep:{s:'abc', c:'red'}}}))")
    check(g, "Sanea valores inválidos", await E(pg, "JSON.stringify(L().entries)") == "{}")
    # archivo real por input
    await pg.set_input_files("#fileIn", files=[{"name": "c.json", "mimeType": "application/json", "buffer": json.dumps(exported).encode()}])
    await pg.wait_for_timeout(300)
    check(g, "Importar desde archivo (input file)", await E(pg, "L().name") == "Exp" and await E(pg, "S.order.length") == n + 2)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_export_respaldo(b):
    g = "12b. Exportar sin capacidad de descarga (claude.ai)"
    ctx = await b.new_context(viewport={"width": 1400, "height": 900})
    await ctx.add_init_script("window.claude={use:async()=>null};")
    pg = await ctx.new_page(); await pg.goto(URL); await pg.wait_for_timeout(400)
    await do_import(pg, SAMPLE_IAT, name="R")
    await pg.click("[data-ptab=layer]"); await pg.click("button[data-act=exportCsv]"); await pg.wait_for_timeout(200)
    check(g, "Muestra la ventana de copiar y pegar", await E(pg, "document.getElementById('dlg').open") and "no está disponible" in await pg.inner_text("#dlg"))
    csv = await pg.input_value("#expTxt")
    check(g, "Contenido completo en la ventana", csv.count("\n") == 11)
    await ctx.close()

async def t_persistencia_local(b):
    g = "13. Persistencia local"
    ctx, pg, errs = await new_page(b)
    await do_import(pg, SAMPLE_IAT, name="Persistida")
    await pg.select_option("#sort", "name"); await pg.check("#dense")
    await pg.reload(); await pg.wait_for_timeout(400)
    check(g, "Capas sobreviven a recargar", await E(pg, "S.order.length") == 2 and await E(pg, "L().name") == "Persistida")
    check(g, "Anotaciones sobreviven", await E(pg, "Object.keys(L().entries).length") == 11)
    check(g, "Preferencias sobreviven", await E(pg, "S.prefs.sort") == "name" and await E(pg, "S.prefs.dense") is True)
    await E(pg, "localStorage.setItem('malapi-navigator-v1','{basura')"); await pg.reload(); await pg.wait_for_timeout(300)
    check(g, "localStorage corrupto: arranca limpio", await E(pg, "S.order.length") == 1 and await pg.locator(".cell").count() == 428)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_remoto(b):
    g = "14. Almacenamiento privado (db simulado) y descargas"
    ctx, pg, errs = await new_page(b, fake=True)
    await pg.wait_for_timeout(1200)
    check(g, "Detecta db + user", "espacio privado" in await pg.inner_text("#saveState"), await pg.inner_text("#saveState"))
    await do_import(pg, SAMPLE_IAT, name="Remota")
    await pg.wait_for_timeout(1300)
    writes = await E(pg, "window.__writes")
    paths = {p for _, p in writes}
    check(g, "Escribe en el subárbol privado data/users/<id>", all(p.startswith("data/users/u123/") for p in paths) and "data/users/u123/workspace" in paths)
    lid = await E(pg, "S.activeId")
    check(g, "Una doc por capa", f"data/users/u123/layer-{lid}" in paths)
    check(g, "Rutas con cantidad par de segmentos", all(len(p.split("/")) % 2 == 0 for p in paths))
    await pg.click("[data-ptab=layer]"); await pg.click("button[data-act=deleteLayer]"); await pg.click("#dlg .dlg-f button.danger")
    await pg.wait_for_timeout(1300)
    check(g, "Eliminar capa borra su doc", ["delete", f"data/users/u123/layer-{lid}"] in await E(pg, "window.__writes"))
    await pg.click(".cell[data-api=Sleep] >> nth=0"); await pg.keyboard.press("3"); await wait_saved(pg)
    # nuevo navegador (sin localStorage) pero mismo db: debe cargar desde db
    await E(pg, "localStorage.clear()")
    await pg.reload()
    await pg.wait_for_function("store.mode==='db' && store.state==='ok' && S.order.length>0 && Object.keys(L().entries).length>0", timeout=10000)
    check(g, "Otro navegador: recupera las capas desde db", await E(pg, "L().entries.Sleep && L().entries.Sleep.s") == 3, await E(pg, "JSON.stringify(Object.keys(S.layers))"))
    await pg.click("[data-ptab=layer]"); await pg.click("button[data-act=exportLayer]"); await pg.wait_for_timeout(300)
    saved = await E(pg, "window.__saved")
    check(g, "Exportar usa la capacidad downloads", len(saved) == 1 and saved[0]["filename"].endswith(".json") and not await E(pg, "document.getElementById('dlg').open"))
    await pg.click("[data-ptab=layer]"); await pg.click("button[data-act=exportCsv]"); await pg.wait_for_timeout(300)
    check(g, "Extensiones permitidas (.json/.csv)", all(s["filename"].rsplit(".", 1)[1] in ("json", "csv") for s in await E(pg, "window.__saved")))
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_movil(b):
    g = "15. Móvil y tema oscuro"
    ctx, pg, errs = await new_page(b, viewport=(390, 844), scheme="dark")
    bg = await E(pg, "getComputedStyle(document.body).backgroundColor")
    check(g, "Tema oscuro aplicado", bg == "rgb(21, 24, 28)", bg)
    check(g, "Sin scroll horizontal del body", await E(pg, "document.documentElement.scrollWidth <= window.innerWidth"))
    check(g, "Panel oculto al inicio", not await E(pg, "document.getElementById('panel').classList.contains('open')"))
    check(g, "Botón Panel visible", await pg.locator(".fab").is_visible())
    await pg.locator(".cell[data-api=VirtualAllocEx]").scroll_into_view_if_needed()
    await pg.click(".cell[data-api=VirtualAllocEx]"); await pg.wait_for_timeout(300)
    check(g, "Tocar una API abre el panel", await E(pg, "document.getElementById('panel').classList.contains('open')"))
    await pg.click(".pclose"); await pg.wait_for_timeout(300)
    check(g, "Cerrar panel", not await E(pg, "document.getElementById('panel').classList.contains('open')"))
    await pg.click(".toolbar button[data-act=importApis]")
    w = await E(pg, "document.getElementById('dlg').getBoundingClientRect().width")
    check(g, "Diálogo cabe en pantalla", w <= 390, w)
    await ctx.close()
    ctx, pg, errs = await new_page(b)
    await E(pg, "document.documentElement.dataset.theme='dark'"); await pg.wait_for_timeout(300)
    check(g, "data-theme=dark fuerza oscuro", await E(pg, "getComputedStyle(document.body).backgroundColor") == "rgb(21, 24, 28)")
    await E(pg, "document.documentElement.dataset.theme='light'"); await pg.wait_for_timeout(300)
    check(g, "data-theme=light fuerza claro", await E(pg, "getComputedStyle(document.body).backgroundColor") == "rgb(236, 238, 233)")
    await ctx.close()

async def t_vista(b):
    g = "17. Modo oscuro y paneles ocultables"
    ctx, pg, errs = await new_page(b, scheme="light")
    BG = "getComputedStyle(document.body).backgroundColor"
    DARK, LIGHT = "rgb(21, 24, 28)", "rgb(236, 238, 233)"
    check(g, "Arranca siguiendo al sistema (claro)", await E(pg, BG) == LIGHT and await E(pg, "document.documentElement.hasAttribute('data-theme')") is False)
    check(g, "Botón indica modo claro", await pg.get_attribute("#ibTheme", "aria-pressed") == "false" and "claro" in await pg.inner_text("#ibTheme"))
    await pg.click("#ibTheme"); await pg.wait_for_timeout(250)
    check(g, "Botón activa modo oscuro", await E(pg, BG) == DARK and await E(pg, "document.documentElement.dataset.theme") == "dark")
    check(g, "Botón refleja estado (aria-pressed=true)", await pg.get_attribute("#ibTheme", "aria-pressed") == "true" and "oscuro" in await pg.inner_text("#ibTheme"))
    await do_import(pg, "VirtualAllocEx\nSleep")
    c = await E(pg, "getComputedStyle(document.querySelector('.cell[data-api=Accept]')).backgroundColor")
    check(g, "Celdas sin anotar usan tokens oscuros", c == "rgb(33, 38, 44)", c)
    await pg.reload(); await pg.wait_for_timeout(400)
    check(g, "Modo oscuro persiste al recargar", await E(pg, BG) == DARK)
    await pg.keyboard.press("d"); await pg.wait_for_timeout(250)
    check(g, "Tecla D alterna el tema", await E(pg, BG) == LIGHT)
    await ctx.close()

    ctx, pg, errs = await new_page(b, scheme="dark")
    check(g, "Con sistema oscuro arranca oscuro y el botón lo indica", await E(pg, BG) == DARK and await pg.get_attribute("#ibTheme", "aria-pressed") == "true")
    await pg.click("#ibTheme"); await pg.wait_for_timeout(250)
    check(g, "Puede forzar claro sobre sistema oscuro", await E(pg, BG) == LIGHT)
    await ctx.close()

    ctx, pg, errs = await new_page(b)
    W = "document.getElementById('mwrap').getBoundingClientRect().width"
    w0 = await E(pg, W)
    await pg.click("#ibPanel")
    check(g, "Ocultar panel lateral", not await pg.locator("#panel").is_visible() and await pg.get_attribute("#ibPanel", "aria-pressed") == "false")
    w1 = await E(pg, W)
    check(g, f"La matriz gana el ancho del panel (+{w1-w0:.0f}px)", w1 - w0 > 300)
    await pg.click(".cell[data-api=VirtualAllocEx]")
    check(g, "Con panel oculto, seleccionar no lo reabre", not await pg.locator("#panel").is_visible())
    check(g, "Leyenda ofrece mostrar panel para editar", await pg.locator("#legend button[data-act=showPanel]").count() == 1)
    await pg.keyboard.press("4")
    check(g, "Atajos siguen funcionando con panel oculto", await E(pg, "L().entries.VirtualAllocEx.s") == 4)
    await pg.click("#legend button[data-act=showPanel]")
    check(g, "Mostrar panel desde la leyenda", await pg.locator("#panel").is_visible() and (await pg.inner_text("#pbody h2")) == "VirtualAllocEx")
    await pg.click(".pclose")
    check(g, "Botón Ocultar dentro del panel", not await pg.locator("#panel").is_visible())
    await pg.keyboard.press("p")
    check(g, "Tecla P muestra el panel", await pg.locator("#panel").is_visible())
    await pg.click("#ibToolbar")
    check(g, "Ocultar barra de herramientas", not await pg.locator(".toolbar").is_visible() and await pg.get_attribute("#ibToolbar", "aria-pressed") == "false")
    await pg.keyboard.press("/")
    check(g, "Tecla / reabre la barra y enfoca la búsqueda", await pg.locator(".toolbar").is_visible() and await E(pg, "document.activeElement.id") == "q")
    await pg.keyboard.press("Escape"); await E(pg, "document.activeElement.blur()")
    await pg.keyboard.press("b")
    check(g, "Tecla B alterna la barra", not await pg.locator(".toolbar").is_visible())
    await pg.click("#ibLegend")
    check(g, "Ocultar leyenda", not await pg.locator("#legend").is_visible() and await pg.get_attribute("#ibLegend", "aria-pressed") == "false")
    await pg.reload(); await pg.wait_for_timeout(400)
    check(g, "Estado de paneles persiste al recargar", not await pg.locator(".toolbar").is_visible() and not await pg.locator("#legend").is_visible() and await pg.locator("#panel").is_visible())
    await pg.keyboard.press("l"); await pg.keyboard.press("b")
    check(g, "Restaurar todo con teclado", await pg.locator(".toolbar").is_visible() and await pg.locator("#legend").is_visible())
    await pg.fill("#q", "pbld")
    check(g, "Letras de atajo no se disparan al escribir", await pg.locator(".toolbar").is_visible() and await pg.locator("#panel").is_visible())
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

    ctx, pg, errs = await new_page(b, viewport=(390, 844))
    check(g, "Móvil: controles del encabezado caben", await E(pg, "document.querySelector('.hctrl').getBoundingClientRect().right <= window.innerWidth") and await E(pg, "document.documentElement.scrollWidth <= window.innerWidth"))
    await pg.click("#ibPanel"); await pg.wait_for_timeout(300)
    check(g, "Móvil: botón Panel abre la hoja inferior", await E(pg, "document.getElementById('panel').classList.contains('open')") and await pg.get_attribute("#ibPanel", "aria-pressed") == "true")
    await pg.click("#ibPanel"); await pg.wait_for_timeout(300)
    check(g, "Móvil: y la cierra", not await E(pg, "document.getElementById('panel').classList.contains('open')"))
    check(g, "Móvil: no altera la preferencia de escritorio", await E(pg, "S.prefs.hidePanel") is False)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_limpiar(b):
    g = "18. Limpiar el canvas"
    ctx, pg, errs = await new_page(b, fake=True)
    await pg.wait_for_timeout(1200)
    await pg.click("#ibClear")
    check(g, "Botón Limpiar abre el menú con 3 opciones", await pg.locator("#dlg .choice button").count() == 3)
    check(g, "Vista limpia: opción deshabilitada", await pg.locator("#dlg .choice button[data-c=view]").is_disabled())
    check(g, "Capa vacía: 'Vaciar' deshabilitado", await pg.locator("#dlg .choice button[data-c=layer]").is_disabled())
    await pg.click("#dlg [data-close]")
    await do_import(pg, SAMPLE_IAT, name="M1")
    await do_import(pg, "Sleep\nBitBlt", name="M2")
    await pg.click(".tab[data-layer] >> nth=1")
    # --- limpiar vista
    await pg.fill("#q", "keylog"); await pg.wait_for_timeout(200)
    await pg.select_option("#dll", "user32"); await pg.select_option("#show", "marked")
    await pg.click("button[data-act=columns]"); await pg.uncheck("#dlg input[value='7']"); await pg.click("#dlg .dlg-f button.primary")
    await pg.click(".cell[data-api=GetAsyncKeyState]")
    await pg.click("#ibClear")
    check(g, "Con filtros activos, 'Limpiar la vista' se habilita", await pg.locator("#dlg .choice button[data-c=view]").is_enabled())
    await pg.click("#dlg .choice button[data-c=view]"); await pg.wait_for_timeout(150)
    check(g, "Limpiar vista: sin búsqueda ni filtros", await pg.input_value("#q") == "" and await pg.input_value("#dll") == "" and await pg.input_value("#show") == "all")
    check(g, "Limpiar vista: sin selección y columnas visibles", await E(pg, "UI.sel.size") == 0 and await pg.locator(".col").count() == 8)
    check(g, "Limpiar vista: 428 celdas y sin atenuado", await pg.locator(".cell").count() == 428 and await pg.locator(".cell.dim").count() == 0)
    check(g, "Limpiar vista: no toca datos", await E(pg, "Object.keys(L().entries).length") == 11 and await E(pg, "S.order.length") == 3)
    # --- vaciar capa activa + deshacer
    await pg.click("#ibClear"); await pg.click("#dlg .choice button[data-c=layer]"); await pg.wait_for_timeout(150)
    check(g, "Vaciar capa activa", await E(pg, "Object.keys(L().entries).length") == 0 and await E(pg, "L().unmatched.length") == 0)
    check(g, "Vaciar conserva nombre y muestra", await E(pg, "L().name") == "M1")
    check(g, "Otras capas intactas", await E(pg, "Object.keys(S.layers[S.order[2]].entries).length") == 2)
    check(g, "Matriz sin celdas coloreadas", await E(pg, "[...document.querySelectorAll('.cell')].every(c=>!c.style.getPropertyValue('--bg-cell'))"))
    check(g, "Aviso con botón Deshacer", await pg.locator("#toast button", has_text="Deshacer").is_visible())
    await pg.click("#toast button"); await pg.wait_for_timeout(150)
    check(g, "Deshacer restaura las 11 anotaciones", await E(pg, "Object.keys(L().entries).length") == 11 and await E(pg, "L().unmatched") == ["GetLastError"])
    check(g, "Deshacer mantiene la capa activa", await E(pg, "L().name") == "M1")
    # Ctrl+Z
    await pg.click("#ibClear"); await pg.click("#dlg .choice button[data-c=layer]"); await pg.wait_for_timeout(100)
    await pg.keyboard.press("Control+z"); await pg.wait_for_timeout(100)
    check(g, "Ctrl+Z deshace el vaciado", await E(pg, "Object.keys(L().entries).length") == 11)
    await pg.keyboard.press("Control+z")
    check(g, "Ctrl+Z sin nada pendiente no hace nada", await E(pg, "Object.keys(L().entries).length") == 11)
    # botón del panel Capa usa el mismo camino
    await pg.click("[data-ptab=layer]"); await pg.click("button[data-act=clearLayer]"); await pg.wait_for_timeout(100)
    check(g, "Vaciar desde el panel Capa también ofrece deshacer", await E(pg, "Object.keys(L().entries).length") == 0 and await pg.locator("#toast button").is_visible())
    await pg.click("#toast button")
    # --- reiniciar + deshacer + db
    before_ids = await E(pg, "S.order.slice()")
    await pg.click("#ibClear"); await pg.click("#dlg .choice button[data-c=all]"); await wait_saved(pg)
    check(g, "Reiniciar deja una sola capa vacía", await E(pg, "S.order.length") == 1 and await E(pg, "Object.keys(L().entries).length") == 0 and await E(pg, "L().name") == "Capa 1")
    check(g, "Reiniciar conserva pesos y preferencias", await E(pg, "S.weights.Injection") == 5 and await E(pg, "S.prefs.sort") == "score")
    writes = await E(pg, "window.__writes")
    check(g, "Reiniciar borra las capas en el almacenamiento privado", all(["delete", f"data/users/u123/layer-{i}"] in writes for i in before_ids))
    new_id = await E(pg, "S.activeId")
    await pg.click("#toast button"); await wait_saved(pg)
    check(g, "Deshacer reinicio restaura las 3 capas", await E(pg, "S.order") == before_ids and await E(pg, "S.order.map(id=>Object.keys(S.layers[id].entries).length)") == [0, 11, 2])
    writes = await E(pg, "window.__writes")
    check(g, "Deshacer reescribe las capas y borra la vacía en db", ["delete", f"data/users/u123/layer-{new_id}"] in writes and ["set", f"data/users/u123/layer-{before_ids[1]}"] in writes[-8:])
    await E(pg, "localStorage.clear()"); await pg.reload()
    await pg.wait_for_function("store.mode==='db' && store.state==='ok' && S.order.length===3", timeout=10000)
    check(g, "Tras deshacer, db queda consistente (recarga desde db)", await E(pg, "S.order.length") == 3 and await E(pg, "Object.keys(S.layers[S.order[1]].entries).length") == 11)
    # sin confirmación de más: cancelar no hace nada
    await pg.click("#ibClear"); await pg.click("#dlg .dlg-f button >> text=Cancelar")
    check(g, "Cancelar no modifica nada", await E(pg, "S.order.length") == 3)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()
    ctx, pg, errs = await new_page(b, viewport=(390, 844))
    check(g, "Móvil: 5 botones del encabezado caben", await E(pg, "document.querySelector('.hctrl').getBoundingClientRect().right <= window.innerWidth") and await E(pg, "document.documentElement.scrollWidth <= window.innerWidth"))
    await ctx.close()

async def t_encabezado(b):
    g = "19. Encabezado adaptable (anchos de iframe reales)"
    HDR = """(()=>{const q=s=>document.querySelector(s), R=e=>e.getBoundingClientRect(), vw=innerWidth;
      const vis=e=>{const r=R(e);return r.width>0&&r.left>=-1&&r.right<=vw+1};
      return {tabsHidden:q('#tabs').scrollWidth>q('#tabs').clientWidth+1,
        namesCut:[...document.querySelectorAll('.tab .tname')].filter(t=>t.scrollWidth>t.clientWidth+1).length,
        buttons:[...document.querySelectorAll('.hctrl .ib')].every(vis), save:vis(q('#saveState')),
        overflow:document.documentElement.scrollWidth>vw+1, tabsW:Math.round(R(q('#tabs')).width)}})()"""
    for w in [948, 1034, 1108, 1204, 1588]:
        ctx, pg, errs = await new_page(b, viewport=(w, 760))
        await E(pg, """(()=>{S.layers[S.order[0]].name='Capa 1';
          for(const n of ['Dropper campaña Q3 variante B','distlib_t64','Superposición ABC']) addLayer(newLayer(n),false);
          renderAll();})()""")
        r = await E(pg, HDR)
        ok = not r["tabsHidden"] and r["namesCut"] == 0 and r["buttons"] and r["save"] and not r["overflow"] and r["tabsW"] > w * 0.9
        check(g, f"{w}px: 4 capas visibles, sin nombres cortados ni desborde", ok, r)
        await ctx.close()
    ctx, pg, errs = await new_page(b, viewport=(1034, 760))
    await E(pg, "for(let i=0;i<14;i++) addLayer(newLayer('Muestra larga número '+i),false); renderAll();")
    r = await E(pg, HDR)
    check(g, "Muchas capas: se desplazan dentro de la fila, sin romper el resto", r["tabsHidden"] and r["buttons"] and r["save"] and not r["overflow"], r)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

async def t_exportacion(b):
    g = "20. Exportación estilo ATT&CK Navigator"
    import zipfile, io, struct, xml.dom.minidom
    import openpyxl
    ctx, pg, errs = await new_page(b, viewport=(1500, 950))
    async def dl_bytes(click):
        async with pg.expect_download() as info:
            await click()
        d = await info.value
        return d.suggested_filename, open(await d.path(), "rb").read()
    await do_import(pg, SAMPLE_IAT, name="Dropper X", sample="invoice.exe", sha="ab" * 32)
    await E(pg, "L().desc='Muestra de prueba'; L().entries.VirtualAllocEx.c='#7a5cc0'; L().entries.Sleep={s:4,c:'#3f8fbf'}; touch(L()); renderAll()")

    # --- menú
    await pg.click("#ibExport")
    check(g, "Menú Exportar con 6 opciones y el filtro de visibles", await pg.locator("#expMenu .mi").count() == 6 and await pg.locator("#expVisible").count() == 1)
    await pg.keyboard.press("Escape")
    check(g, "Esc cierra el menú", await pg.locator("#expMenu").count() == 0)
    await pg.click("#ibExport"); await pg.mouse.click(300, 600)
    check(g, "Clic afuera cierra el menú", await pg.locator("#expMenu").count() == 0)

    # --- leyenda de colores
    await pg.click("[data-ptab=layer]")
    await pg.click("button[data-act=lgUsed]")
    check(g, "Leyenda: 'Agregar colores usados' toma los 2 colores manuales", await E(pg, "L().legend.map(x=>x.color).sort().join()") == "#3f8fbf,#7a5cc0")
    await pg.locator('[data-lg="label"]').nth(0).fill("Confirmado en sandbox")
    check(g, "Leyenda: la etiqueta se guarda al escribir", "Confirmado en sandbox" in await E(pg, "L().legend.map(x=>x.label).join('|')"))
    check(g, "Leyenda visible bajo la barra de herramientas", "Confirmado en sandbox" in await pg.inner_text("#legend"))
    await pg.click("button[data-act=lgAdd]")
    check(g, "Leyenda: agregar elemento con color libre", await E(pg, "L().legend.length") == 3)
    await pg.click("button[data-act=lgDel][data-v='2']")
    check(g, "Leyenda: quitar elemento", await E(pg, "L().legend.length") == 2)

    # --- SVG: estructura
    P = "(()=>{const s=renderLayerSVG(L(),Object.assign(svgCfg(),%s)); const d=new DOMParser().parseFromString(s,'image/svg+xml'); const r=d.documentElement;" \
        " return {err:!!d.querySelector('parsererror'), w:r.getAttribute('width'), h:r.getAttribute('height'), vb:r.getAttribute('viewBox'), apis:+r.getAttribute('data-apis')," \
        " cells:d.querySelectorAll('g.xcell').length, sec:[...d.querySelectorAll('[data-section]')].map(x=>x.dataset.section), font:r.getAttribute('font-family')," \
        " stroke:(d.querySelector('g.xcell rect')||{getAttribute:()=>null}).getAttribute('stroke'), bg:d.querySelector('svg > rect').getAttribute('fill')}})()"
    S1 = lambda cfg: E(pg, P % json.dumps(cfg))
    r = await S1({})
    check(g, "SVG válido (XML sin errores)", not r["err"], r)
    check(g, "Tamaño físico 11 × 8,5 in con viewBox en px", r["w"] == "11in" and r["h"] == "8.5in" and r["vb"] == "0 0 1056 816", r)
    check(g, "Vista actual: 428 celdas", r["cells"] == 428 == r["apis"])
    check(g, "Encabezado con Capa, Muestra, Filtros y Leyenda", r["sec"] == ["about", "sample", "filters", "legend", "table"], r["sec"])
    ann = await E(pg, "API.filter(a=>isAnnotated(L().entries[a.name])).reduce((s,a)=>s+a.cats.length,0)")
    r = await S1({"include": "annotated"})
    check(g, f"Solo anotadas: {ann} celdas", r["cells"] == ann, r["cells"])
    await pg.select_option("#dll", "user32")
    r1 = await S1({"include": "view"}); r2_ = await S1({"include": "all"})
    check(g, "Según la vista respeta el filtro de DLL; 'todas' lo ignora", r1["cells"] < 100 and r2_["cells"] == 428, (r1["cells"], r2_["cells"]))
    await pg.select_option("#dll", "")
    r = await S1({"showHeader": False})
    check(g, "Sin encabezado: solo la tabla", r["sec"] == ["table"], r["sec"])
    r = await S1({"showAbout": False, "showFilters": False})
    check(g, "Secciones del encabezado se ocultan por separado", r["sec"] == ["sample", "legend", "table"], r["sec"])
    r = await S1({"legendDocked": False, "legendX": 7, "legendY": 5})
    fx = await E(pg, "(()=>{const d=new DOMParser().parseFromString(renderLayerSVG(L(),Object.assign(svgCfg(),{legendDocked:false,legendX:7,legendY:5})),'image/svg+xml'); return +d.querySelector('[data-section=legend-floating] rect').getAttribute('x')})()")
    check(g, "Leyenda flotante en X=7 in, Y=5 in (fuera del encabezado)", "legend-floating" in r["sec"] and "legend" not in r["sec"] and abs(fx - 7 * 96) < 1, (r["sec"], fx))
    r = await S1({"showHeader": False, "legendDocked": False})
    check(g, "Leyenda flotante se muestra aunque no haya encabezado (como Navigator)", r["sec"] == ["table", "legend-floating"], r["sec"])
    r = await S1({"showLegend": False, "showGradient": False})
    check(g, "Sin leyenda ni gradiente: no hay recuadro de leyenda", "legend" not in r["sec"], r["sec"])
    r = await S1({"font": "monospace", "borderColor": "#ff0000", "theme": "dark"})
    check(g, "Fuente, borde de celdas y tema oscuro aplicados", "Consolas" in r["font"] and r["stroke"] == "#ff0000" and r["bg"] == "#1f2328", r)
    r = await S1({"unit": "cm", "width": 29.7, "height": 21})
    check(g, "Unidad cm: 29,7 × 21 cm", r["w"] == "29.7cm" and r["vb"].startswith("0 0 1122.5"), r)
    r = await S1({"unit": "px", "width": 1920, "height": 1080})
    check(g, "Unidad px: 1920 × 1080", r["w"] == "1920" and r["vb"] == "0 0 1920 1080", r)
    r = await S1({"cellText": "none"})
    none_txt = await E(pg, "(()=>{const d=new DOMParser().parseFromString(renderLayerSVG(L(),Object.assign(svgCfg(),{cellText:'none'})),'image/svg+xml'); return d.querySelectorAll('g.xcell text').length})()")
    check(g, "Vista mini: celdas sin texto", r["cells"] == 428 and none_txt == 0)

    # --- el texto entra en cada celda (medido en el navegador)
    FIT = """(cfg)=>{const host=document.createElement('div'); host.style.cssText='position:fixed;left:0;top:0;width:1400px;opacity:0;pointer-events:none';
      host.innerHTML=renderLayerSVG(L(),Object.assign(svgCfg(),cfg)); document.body.appendChild(host);
      const svg=host.querySelector('svg'); svg.setAttribute('width',svg.viewBox.baseVal.width); svg.setAttribute('height',svg.viewBox.baseVal.height);
      let bad=[], n=0; for(const c of svg.querySelectorAll('g.xcell')){ const r=c.querySelector('rect').getBBox();
        for(const t of c.querySelectorAll('text')){ n++; const b=t.getBBox(); if(b.x<r.x-0.5||b.x+b.width>r.x+r.width+0.5||b.y<r.y-0.5||b.y+b.height>r.y+r.height+0.5) bad.push(c.dataset.api); } }
      host.remove(); return {n, bad:[...new Set(bad)].slice(0,8)};}"""
    for cfg, lbl in [({"include": "annotated", "cellText": "name_score"}, "solo anotadas con score"), ({"include": "all"}, "catálogo completo"),
                     ({"include": "all", "font": "serif", "unit": "cm", "width": 21, "height": 29.7}, "A4 vertical, serif"),
                     ({"include": "annotated", "font": "monospace", "unit": "px", "width": 800, "height": 500}, "800 × 500 px, monospace")]:
        r = await pg.evaluate(FIT, cfg)
        check(g, f"Texto dentro de su celda: {lbl} ({r['n']} textos)", r["n"] > 0 and not r["bad"], r["bad"])

    # --- vista de render: controles
    await pg.click("#ibExport"); await pg.click("#expMenu [data-exp=svg]"); await pg.wait_for_timeout(300)
    check(g, "Vista de render abre con vista previa", await pg.locator("#rvPrev svg").count() == 1)
    fam = await E(pg, "getComputedStyle(document.querySelector('#rvPrev g.xcell text')).fontFamily")
    check(g, "La vista previa usa la fuente elegida, no la de la matriz", "Arial" in fam and "Mono" not in fam, fam)
    sel_before = await E(pg, "[...UI.sel].join()")
    await pg.locator("#rvPrev g.xcell").nth(3).click()
    check(g, "Clic en la vista previa no cambia la selección de la matriz", await E(pg, "[...UI.sel].join()") == sel_before and await E(pg, "document.getElementById('dlg').open"))
    await pg.click("#rvUnit [data-u=cm]"); await pg.wait_for_timeout(50)
    check(g, "Cambiar a cm convierte las medidas (11 in → 27,94 cm)", await pg.input_value("#rvW") == "27.94" and await pg.input_value("#rvH") == "21.59")
    await pg.click("#rvUnit [data-u=px]"); await pg.wait_for_timeout(50)
    check(g, "Cambiar a px (→ 1056 × 816)", await pg.input_value("#rvW") == "1056" and await pg.input_value("#rvH") == "816")
    await pg.click("#rvRotate"); await pg.wait_for_timeout(50)
    check(g, "Rotar intercambia ancho y alto", await pg.input_value("#rvW") == "816")
    await pg.select_option("#rvPreset", "a4"); await pg.wait_for_timeout(50)
    check(g, "Preset A4 respeta la orientación vertical", await E(pg, "[svgCfg().unit,svgCfg().width,svgCfg().height].join()") == "cm,21,29.7")
    await pg.fill("#rvW", "20"); await pg.wait_for_timeout(250)
    check(g, "Editar ancho actualiza la vista previa", (await pg.get_attribute("#rvPrev svg", "width")) == "20cm")
    await pg.uncheck("#rvShowHeader"); await pg.wait_for_timeout(250)
    check(g, "Sin encabezado: se deshabilitan sus sub-opciones", await pg.locator("#rvShowAbout").is_disabled() and await pg.locator("#rvHH").is_disabled())
    await pg.uncheck("#rvDocked"); await pg.wait_for_timeout(250)
    check(g, "Desacoplar muestra posición y tamaño de la leyenda", await pg.locator("#rvLX").is_visible() and await pg.locator("#rvLW").is_visible())
    await pg.click("#dlg .dlg-f button:has-text('Restablecer')"); await pg.wait_for_timeout(300)
    check(g, "Restablecer vuelve a Carta 11 × 8,5 in", await E(pg, "[svgCfg().unit,svgCfg().width,svgCfg().height,svgCfg().showHeader].join()") == "in,11,8.5,true")
    # descargas
    fn, data = await dl_bytes(lambda: pg.click("#dlg .dlg-f button:has-text('Descargar SVG')"))
    check(g, "Descargar SVG", fn == "malapi-capa-Dropper_X.svg" and data.startswith(b"<svg") and b"</svg>" in data, fn)
    fn, data = await dl_bytes(lambda: pg.click("#dlg .dlg-f button:has-text('Descargar PNG')"))
    w, h = struct.unpack(">II", data[16:24])
    check(g, f"Descargar PNG 2× (2112 × 1632 px; obtuvo {w} × {h})", data[:8] == b"\x89PNG\r\n\x1a\n" and (w, h) == (2112, 1632))
    await pg.click("#rvScale [data-s='3']"); await pg.wait_for_timeout(100)
    check(g, "Resolución 3× informa el tamaño en px", "3168 × 2448" in await pg.inner_text("#rvPxInfo"))
    fn, data = await dl_bytes(lambda: pg.click("#dlg .dlg-f button:has-text('Descargar PNG')"))
    check(g, "PNG 3× (3168 × 2448)", struct.unpack(">II", data[16:24]) == (3168, 2448))
    await pg.click("#dlg [data-close]")

    # --- Excel
    await E(pg, "S.prefs.hiddenCols=[4]; S.prefs.colsByWeight=true; renderAll()")
    await pg.click("#ibExport")
    fn, data = await dl_bytes(lambda: pg.click("#expMenu [data-exp=xlsx1]"))
    z = zipfile.ZipFile(io.BytesIO(data))
    check(g, "XLSX: ZIP íntegro (CRC)", z.testzip() is None and fn == "malapi-capa-Dropper_X.xlsx")
    bad_xml = []
    for n in z.namelist():
        try: xml.dom.minidom.parseString(z.read(n))
        except Exception as ex: bad_xml.append((n, str(ex)[:60]))
    check(g, "XLSX: todas las partes son XML válido", not bad_xml and "[Content_Types].xml" in z.namelist(), bad_xml)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    check(g, "XLSX: hoja matriz y hoja detalle", wb.sheetnames == ["Dropper X", "Dropper X (detalle)"], wb.sheetnames)
    ws = wb["Dropper X"]
    heads = [c.value for c in ws[1]]
    expected_heads = await E(pg, "columns(L()).map(c=>c.cat)")
    check(g, "XLSX: columnas en el orden de la vista y sin las ocultas", heads == expected_heads and "Internet" not in heads, heads)
    colors = {}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value and c.fill and c.fill.fill_type == "solid": colors[c.value] = c.fill.fgColor.rgb[-6:].lower()
    exp_sleep = "3f8fbf"; exp_vae = "7a5cc0"
    exp_grad = (await E(pg, "cellColor(L(),L().entries.OpenProcess)"))[1:]
    check(g, "XLSX: colores manuales y de gradiente iguales a la matriz", colors.get("Sleep") == exp_sleep and colors.get("VirtualAllocEx") == exp_vae and colors.get("OpenProcess") == exp_grad, (colors.get("Sleep"), colors.get("VirtualAllocEx"), colors.get("OpenProcess"), exp_grad))
    un = [c.value for row in ws.iter_rows(min_row=2) for c in row if c.value == "EnumSystemLocalesA"]
    check(g, "XLSX: APIs sin anotar presentes sin relleno", len(un) >= 1)
    det = wb["Dropper X (detalle)"]
    rows = list(det.iter_rows(min_row=2, values_only=True))
    check(g, "XLSX detalle: anotadas visibles, ordenadas por score", rows[0][0] == "GetAsyncKeyState" and rows[0][4] == 7 and all(r[0] != "InternetOpenA" for r in rows), rows[:2])
    check(g, "XLSX: panel fijo en la fila de encabezados", ws.freeze_panes == "A2")
    await E(pg, "(()=>{const x=newLayer('Capa [con]: *símbolos?/'); x.entries={Sleep:{s:2}}; addLayer(x,false); const y=newLayer('Dropper X'); addLayer(y,false); renderAll();})()")
    await pg.click("#ibExport")
    fn, data = await dl_bytes(lambda: pg.click("#expMenu [data-exp=xlsxAll]"))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    names = wb.sheetnames
    check(g, "XLSX todas las capas: 2 hojas por capa", len(names) == 2 * await E(pg, "S.order.length"), names)
    check(g, "XLSX: nombres de hoja saneados, únicos y ≤ 31 caracteres", all(len(n) <= 31 and not any(ch in n for ch in "[]:*?/\\") for n in names) and len(set(n.lower() for n in names)) == len(names), names)

    # --- solo anotaciones visibles (JSON/CSV)
    await E(pg, "S.activeId=S.order[1]; S.prefs.hiddenCols=[7]; renderAll()")
    await pg.click("#ibExport"); await pg.check("#expVisible")
    fn, data = await dl_bytes(lambda: pg.click("#expMenu [data-exp=json1]"))
    apis = [e["api"] for e in json.loads(data)["entries"]]
    check(g, "JSON 'solo visibles' omite APIs de columnas ocultas", "RegOpenKeyExA" not in apis and "VirtualAllocEx" in apis, apis)
    check(g, "JSON incluye legendItems (formato Navigator)", json.loads(data)["legendItems"][0]["label"] == "Confirmado en sandbox" or json.loads(data)["legendItems"][1]["label"] == "Confirmado en sandbox")
    await pg.click("#ibExport")
    fn, data = await dl_bytes(lambda: pg.click("#expMenu [data-exp=csv1]"))
    check(g, "CSV 'solo visibles' también filtra", b"RegOpenKeyExA" not in data and b"VirtualAllocEx" in data)
    await pg.click("#ibExport"); await pg.uncheck("#expVisible")
    fn, data = await dl_bytes(lambda: pg.click("#expMenu [data-exp=json1]"))
    check(g, "Sin el filtro exporta todas las anotaciones", "RegOpenKeyExA" in [e["api"] for e in json.loads(data)["entries"]])

    # --- importar legendItems de Navigator y herencia en superposición
    nav = {"name": "Desde Navigator", "techniques": [{"techniqueID": "Sleep", "score": 1, "color": "#e60d0d"}], "legendItems": [{"label": "Visto en campo", "color": "#E60D0D"}]}
    await E(pg, f"importJsonText({json.dumps(json.dumps(nav))})")
    check(g, "Importa legendItems de ATT&CK Navigator", await E(pg, "JSON.stringify(L().legend)") == '[{"color":"#e60d0d","label":"Visto en campo"}]')
    await E(pg, "S.activeId=S.order[1]; renderAll()")
    rec = {"sources": [["a", "__A__"], ["b", "__B__"]], "expr": "a + b", "colorFrom": "__A__", "comments": False, "colorOnly": 1}
    ids = await E(pg, "S.order.slice()")
    rec_js = json.dumps(rec).replace("__A__", ids[1]).replace("__B__", ids[-1])
    await E(pg, f"(()=>{{const r={rec_js}; const res=overlayCompute(r); const l=newLayer('Sup'); applyOverlayResult(l,r,res); addLayer(l); renderAll();}})()")
    check(g, "Superposición hereda la leyenda de la capa de color", await E(pg, "L().legend.length") == 2)

    # --- persistencia
    await E(pg, "S.prefs.svg=Object.assign(svgCfg(),{unit:'cm',width:30,theme:'dark'}); touch(null)")
    await pg.reload(); await pg.wait_for_timeout(400)
    check(g, "Configuración de imagen y leyendas sobreviven a recargar", await E(pg, "[svgCfg().unit,svgCfg().width,svgCfg().theme].join()") == "cm,30,dark" and await E(pg, "S.layers[S.order[1]].legend.length") == 2)
    check(g, "Sin errores JS", not errs, errs)
    await ctx.close()

    # --- claude.ai: la capacidad de descarga recibe los archivos binarios
    ctx, pg, errs = await new_page(b, fake=True)
    await pg.wait_for_timeout(800)
    await do_import(pg, SAMPLE_IAT, name="Claude")
    await E(pg, "openRenderView()"); await pg.wait_for_timeout(300)
    await pg.click("#dlg .dlg-f button:has-text('Descargar SVG')"); await pg.wait_for_timeout(200)
    await pg.click("#dlg .dlg-f button:has-text('Descargar PNG')"); await pg.wait_for_timeout(800)
    await pg.click("#dlg [data-close]")
    await pg.click("#ibExport"); await pg.click("#expMenu [data-exp=xlsx1]"); await pg.wait_for_timeout(300)
    saved = await E(pg, "window.__saved.map(s=>s.filename)")
    check(g, "claude.ai: SVG, PNG y XLSX pasan por la capacidad de descarga", saved == ["malapi-capa-Claude.svg", "malapi-capa-Claude.png", "malapi-capa-Claude.xlsx"], saved)
    check(g, "Sin errores JS (claude.ai)", not errs, errs)
    await ctx.close()

async def t_rendimiento(b):
    g = "16. Rendimiento"
    ctx, pg, errs = await new_page(b)
    await pg.click(".tab.add"); await pg.click(".choice button[data-c=reference]")
    ms = await E(pg, "(()=>{const t=performance.now();for(let i=0;i<20;i++)renderMatrix();return (performance.now()-t)/20})()")
    check(g, f"Render de matriz completa < 50 ms (medido {ms:.1f} ms)", ms < 50, ms)
    size = await E(pg, "JSON.stringify(L()).length")
    check(g, f"Capa llena entra en una doc de db (<256 KB; {size/1024:.1f} KB)", size < 256 * 1024)
    await ctx.close()

async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()
        for t in [t_carga, t_matching, t_parser, t_import_ui, t_seleccion_edicion, t_filtros, t_capa, t_expr, t_overlay, t_overlay_bordes,
                  t_perfil, t_pesos, t_export_import, t_export_respaldo, t_persistencia_local, t_remoto, t_movil, t_vista, t_limpiar, t_encabezado, t_exportacion, t_rendimiento]:
            try:
                await t(b)
            except Exception as e:
                RESULTS.append((t.__name__, "EXCEPCIÓN", False, traceback.format_exc(limit=3)))
        await b.close()
    cur = None; fails = 0
    for grp, name, ok, det in RESULTS:
        if grp != cur: print(f"\n{grp}"); cur = grp
        print(f"  {'PASS' if ok else 'FAIL'}  {name}" + ("" if ok else f"   -> {det}"))
        fails += (not ok)
    print(f"\nTOTAL: {len(RESULTS)-fails}/{len(RESULTS)} OK, {fails} fallas")
    sys.exit(1 if fails else 0)

asyncio.run(main())
